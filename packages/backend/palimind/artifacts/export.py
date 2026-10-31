"""Server-side artifact export fallbacks.

The UI performs most exports client-side (jsPDF/DOCX/PPTX/canvas). These
helpers keep the API contract working for clients that prefer a server-side
download, and they persist generated binaries in the owning project.
"""

from __future__ import annotations

import html as html_lib
import io
import re
from pathlib import Path
from typing import Any

_FORMATS = {
    "md": ("md", "text/markdown; charset=utf-8"),
    "html": ("html", "text/html; charset=utf-8"),
    "pdf": ("pdf", "application/pdf"),
    "docx": ("docx", "application/vnd.openxmlformats-officedocument.wordprocessingml.document"),
    "pptx": ("pptx", "application/vnd.openxmlformats-officedocument.presentationml.presentation"),
    "png": ("png", "image/png"),
}


def _plain_text(content: str) -> str:
    text = re.sub(r"```[^\n]*\n?", "", content)
    text = text.replace("```", "")
    text = re.sub(r"^#{1,6}\s+", "", text, flags=re.MULTILINE)
    return text


def _html_document(title: str, content: str, artifact_type: str) -> str:
    if artifact_type == "html" and "<html" in content.lower():
        return content
    body = (
        content
        if artifact_type in {"html", "svg"}
        else f"<pre>{html_lib.escape(content)}</pre>"
    )
    return f"""<!doctype html>
<html lang="en"><head><meta charset="utf-8" />
<meta name="viewport" content="width=device-width, initial-scale=1" />
<title>{html_lib.escape(title)}</title>
<style>
body {{ margin: 0; padding: 40px; font: 15px/1.65 system-ui, sans-serif; background: #fff; color: #18181b; }}
main {{ max-width: 860px; margin: 0 auto; }}
pre {{ white-space: pre-wrap; overflow-wrap: anywhere; background: #f4f4f5; padding: 16px; border-radius: 10px; }}
svg, img {{ max-width: 100%; height: auto; }}
</style></head><body><main>{body}</main></body></html>"""


def _pdf_bytes(title: str, content: str, artifact_type: str) -> bytes:
    try:
        from weasyprint import HTML

        return HTML(string=_html_document(title, content, artifact_type)).write_pdf()
    except Exception:
        return _minimal_pdf(title, _plain_text(content))


def _minimal_pdf(title: str, text: str) -> bytes:
    """Dependency-free text PDF fallback (Helvetica, A4, simple line wrap)."""
    width, height = 595, 842
    margin, line_height = 48, 15
    max_chars = 86
    lines: list[str] = [title, ""]
    for raw in text.splitlines():
        if not raw:
            lines.append("")
            continue
        while len(raw) > max_chars:
            split_at = raw.rfind(" ", 0, max_chars)
            if split_at <= 0:
                split_at = max_chars
            lines.append(raw[:split_at])
            raw = raw[split_at:].lstrip()
        lines.append(raw)

    pages: list[list[str]] = []
    current: list[str] = []
    capacity = (height - margin * 2) // line_height
    for line in lines:
        if len(current) >= capacity:
            pages.append(current)
            current = []
        current.append(line)
    if current:
        pages.append(current)
    if not pages:
        pages = [[""]]

    def escape(value: str) -> str:
        return value.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")

    objects: list[bytes] = []
    objects.append(b"<< /Type /Catalog /Pages 2 0 R >>")
    page_ids = [4 + i * 2 for i in range(len(pages))]
    kids = " ".join(f"{pid} 0 R" for pid in page_ids)
    objects.append(f"<< /Type /Pages /Kids [{kids}] /Count {len(pages)} >>".encode())
    objects.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")
    for index, page in enumerate(pages):
        page_id = 4 + index * 2
        content_id = page_id + 1
        lines_pdf = []
        y = height - margin
        for line in page:
            lines_pdf.append(
                f"BT /F1 10 Tf 1 0 0 1 {margin} {y} Tm ({escape(line)}) Tj ET"
            )
            y -= line_height
        stream = "\n".join(lines_pdf).encode("latin-1", "replace")
        objects.append(
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 {width} {height}] "
            f"/Resources << /Font << /F1 3 0 R >> >> /Contents {content_id} 0 R >>".encode()
        )
        objects.append(b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream")

    out = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for i, obj in enumerate(objects, start=1):
        offsets.append(len(out))
        out.extend(f"{i} 0 obj\n".encode())
        out.extend(obj)
        out.extend(b"\nendobj\n")
    xref = len(out)
    out.extend(f"xref\n0 {len(objects) + 1}\n".encode())
    out.extend(b"0000000000 65535 f \n")
    for offset in offsets[1:]:
        out.extend(f"{offset:010d} 00000 n \n".encode())
    out.extend(
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    )
    return bytes(out)


def _docx_bytes(title: str, content: str) -> bytes:
    from docx import Document

    document = Document()
    document.add_heading(title or "Artifact", level=0)
    in_code = False
    for raw in content.splitlines():
        line = raw.rstrip()
        if line.strip().startswith("```"):
            in_code = not in_code
            continue
        if in_code:
            paragraph = document.add_paragraph()
            run = paragraph.add_run(line or " ")
            run.font.name = "Courier New"
            continue
        heading = re.match(r"^(#{1,4})\s+(.*)$", line)
        if heading:
            document.add_heading(heading.group(2), level=min(4, len(heading.group(1))))
            continue
        bullet = re.match(r"^\s*[-*]\s+(.*)$", line)
        if bullet:
            document.add_paragraph(bullet.group(1), style="List Bullet")
            continue
        document.add_paragraph(line)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _slide_specs(title: str, content: str, artifact_type: str) -> list[tuple[str, list[str]]]:
    slides: list[tuple[str, list[str]]] = []
    current_title = ""
    bullets: list[str] = []

    def flush() -> None:
        nonlocal current_title, bullets
        if current_title or bullets:
            slides.append((current_title or title or "Slide", bullets))
        current_title, bullets = "", []

    if artifact_type != "slide-outline":
        current_title = title or "Artifact"
        for line in content.splitlines():
            text = re.sub(r"^#{1,6}\s*", "", line).strip()
            text = re.sub(r"^[-*]\s+", "", text)
            if text:
                bullets.append(text)
        flush()
        return slides or [(title or "Artifact", [])]

    for raw in content.splitlines():
        line = raw.strip()
        if not line:
            continue
        if re.match(r"^-{3,}$", line):
            flush()
            continue
        heading = re.match(r"^#{1,4}\s+(.+)$", line) or re.match(
            r"^slide\s*\d+\s*[:.-]\s*(.+)$", line, re.IGNORECASE
        )
        if heading:
            flush()
            current_title = heading.group(1)
            continue
        bullets.append(re.sub(r"^\s*[-*•]\s+", "", line))
    flush()
    return slides or [(title or "Artifact", [])]


def _pptx_bytes(title: str, content: str, artifact_type: str) -> bytes:
    from pptx import Presentation
    from pptx.util import Inches, Pt

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    cover = prs.slides.add_slide(blank)
    box = cover.shapes.add_textbox(Inches(0.7), Inches(2.2), Inches(11.9), Inches(1.4))
    frame = box.text_frame
    frame.text = title or "Artifact"
    frame.paragraphs[0].font.size = Pt(36)
    frame.paragraphs[0].font.bold = True
    subtitle = cover.shapes.add_textbox(Inches(0.7), Inches(3.8), Inches(11.9), Inches(0.7))
    subtitle.text_frame.text = "Generated with Palimind"

    for slide_title, bullets in _slide_specs(title, content, artifact_type):
        slide = prs.slides.add_slide(blank)
        heading = slide.shapes.add_textbox(Inches(0.6), Inches(0.45), Inches(12.1), Inches(0.9))
        heading.text_frame.text = slide_title
        heading.text_frame.paragraphs[0].font.size = Pt(28)
        heading.text_frame.paragraphs[0].font.bold = True
        if bullets:
            body = slide.shapes.add_textbox(Inches(0.8), Inches(1.5), Inches(11.7), Inches(5.3))
            tf = body.text_frame
            tf.word_wrap = True
            for index, bullet in enumerate(bullets[:12]):
                paragraph = tf.paragraphs[0] if index == 0 else tf.add_paragraph()
                paragraph.text = f"• {bullet}"
                paragraph.font.size = Pt(18)

    buffer = io.BytesIO()
    prs.save(buffer)
    return buffer.getvalue()


def _png_bytes(title: str, content: str) -> bytes:
    from PIL import Image, ImageDraw

    width, height = 1600, 2000
    image = Image.new("RGB", (width, height), "#ffffff")
    draw = ImageDraw.Draw(image)
    draw.text((70, 60), title or "Artifact", fill="#18181b")
    y = 130
    for line in _plain_text(content).splitlines():
        if y > height - 80:
            break
        draw.text((70, y), line[:140], fill="#3f3f46")
        y += 26
    buffer = io.BytesIO()
    image.save(buffer, format="PNG")
    return buffer.getvalue()


def export_filename(title: str, extension: str) -> str:
    base = re.sub(r"[^\w.-]+", "_", title or "artifact").strip("._") or "artifact"
    return f"{base}.{extension}"


_ALIASES = {
    "markdown": "md",
    "htm": "html",
    "word": "docx",
    "powerpoint": "pptx",
    "image": "png",
}


def normalize_format(fmt: str) -> str:
    normalized = fmt.lower().lstrip(".")
    return _ALIASES.get(normalized, normalized)


def export_bytes(artifact: dict[str, Any], fmt: str) -> tuple[bytes, str, str]:
    """Return ``(bytes, extension, media_type)`` for an artifact export."""
    normalized = normalize_format(fmt)
    if normalized not in _FORMATS:
        raise ValueError(f"Unsupported export format: {fmt}")
    artifact_type = str(artifact.get("type") or "code")
    title = str(artifact.get("title") or "Artifact")
    content = str(artifact.get("content") or "")

    if normalized == "md":
        extension, media = _FORMATS[normalized]
        if artifact_type in {"markdown", "document", "table", "slide-outline"}:
            text = content
        else:
            text = f"# {title}\n\n```{artifact.get('language') or artifact_type}\n{content}\n```\n"
        return text.encode("utf-8"), extension, media
    if normalized == "html":
        extension, media = _FORMATS[normalized]
        return _html_document(title, content, artifact_type).encode("utf-8"), extension, media
    if normalized == "pdf":
        extension, media = _FORMATS[normalized]
        return _pdf_bytes(title, content, artifact_type), extension, media
    if normalized == "docx":
        extension, media = _FORMATS[normalized]
        return _docx_bytes(title, content), extension, media
    if normalized == "pptx":
        extension, media = _FORMATS[normalized]
        return _pptx_bytes(title, content, artifact_type), extension, media
    extension, media = _FORMATS["png"]
    return _png_bytes(title, content), extension, media


def store_export(
    *,
    artifact_id: str,
    fmt: str,
    data: bytes,
    project_id: str = "",
) -> Path:
    """Persist an exported binary in the project (or global) artifact folder."""
    from palimind.projects.store import project_dir
    from palimind.security.tokens import palimind_home

    extension = _FORMATS[normalize_format(fmt)][0]
    if project_id:
        root = project_dir(project_id) / "artifacts" / "exports"
    else:
        root = palimind_home() / "artifacts" / "exports"
    root.mkdir(parents=True, exist_ok=True)
    safe_id = re.sub(r"[^A-Za-z0-9_.-]+", "-", artifact_id)[:80] or "artifact"
    path = root / f"{safe_id}.{extension}"
    path.write_bytes(data)
    return path
